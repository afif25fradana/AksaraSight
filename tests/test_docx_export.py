"""Unit tests for DOCX export converter (core/docx_export.py)."""

import io
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from markdown_it.token import Token
from markdown_it.tree import SyntaxTreeNode

from core.docx_export import (
    _HTMLTableExtractor,
    _append_table_truncation_note,
    _render_ast_node,
    _render_parsed_html_table,
    _render_table_node,
    build_docx,
    export_to_docx_bytes,
    render_markdown_page,
)
from core.models import JobStatus, OCRResult, PageResult


# Page Breaks

def test_page_breaks_between_pages_only(tmp_path: Path) -> None:
    """Verify literal page breaks are inserted strictly between pages and never after the last page."""
    result = OCRResult(
        file_path="multipage.pdf",
        pages=[
            PageResult(page_num=1, markdown="# Page One\nFirst page content.", status=JobStatus.SUCCESS),
            PageResult(page_num=2, markdown="# Page Two\nSecond page content.", status=JobStatus.SUCCESS),
            PageResult(page_num=3, markdown="# Page Three\nThird page content.", status=JobStatus.SUCCESS),
        ],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "multipage_breaks.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Count page breaks across all runs in all paragraphs
    # In python-docx/OOXML, a page break is represented as a <w:br w:type="page"/> element within a run
    page_break_count = 0
    for p in paragraphs:
        for r in p.runs:
            if '<w:br w:type="page"/>' in r._r.xml:
                page_break_count += 1

    # Between 3 pages, there must be exactly 2 page breaks
    assert page_break_count == 2, f"Expected exactly 2 page breaks for 3 pages, found {page_break_count}"

    # Verify the last paragraph does not have a trailing page break
    last_paragraph = paragraphs[-1]
    for r in last_paragraph.runs:
        assert '<w:br w:type="page"/>' not in r._r.xml


def test_internal_horizontal_rule_is_not_page_break(tmp_path: Path) -> None:
    """Verify markdown '---' inside a single page is rendered as visual divider, not a page break."""
    result = OCRResult(
        file_path="single_page_divider.pdf",
        pages=[
            PageResult(
                page_num=1,
                markdown="Top Section\n\n---\n\nBottom Section",
                status=JobStatus.SUCCESS,
            ),
        ],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "single_hr.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Verify zero page breaks exist
    page_breaks = [
        r for p in paragraphs for r in p.runs
        if '<w:br w:type="page"/>' in r._r.xml
    ]
    assert len(page_breaks) == 0, "Internal horizontal rule '---' must not produce a page break"

    # Verify visual divider exists (rendered with divider characters)
    divider_found = any("―" in p.text for p in paragraphs)
    assert divider_found, "Internal horizontal rule should render as visual divider paragraph"


# GFM Tables & OOXML Round-Trip Tests

def test_table_ooxml_header_repeat_and_cant_split(tmp_path: Path) -> None:
    """Verify OOXML w:tblHeader and w:cantSplit attributes on raw XML round-trip."""
    table_md = (
        "| School Region | Total Enrolled | Pass Rate |\n"
        "|:---|:---:|---:|\n"
        "| Region 1 | 1,250 | 94.5% |\n"
        "| Region 2 | 890 | 88.2% |\n"
    )
    result = OCRResult(
        file_path="table.pdf",
        pages=[PageResult(page_num=1, markdown=table_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "ooxml_table.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    assert len(reopened.tables) == 1
    table = reopened.tables[0]

    # Verify style
    assert table.style is not None
    assert table.style.name == "Table Grid"

    # Row 0: Header row must have tblHeader and cantSplit in raw XML
    header_tr_xml = table.rows[0]._tr.xml
    assert "<w:tblHeader/>" in header_tr_xml or "<w:tblHeader" in header_tr_xml, (
        "Header row missing <w:tblHeader/> repeat element in OOXML"
    )
    assert "<w:cantSplit/>" in header_tr_xml or "<w:cantSplit" in header_tr_xml, (
        "Header row missing <w:cantSplit/> element in OOXML"
    )

    # Row 1: Data row must have cantSplit, but NOT tblHeader
    data_tr_xml = table.rows[1]._tr.xml
    assert "<w:cantSplit/>" in data_tr_xml or "<w:cantSplit" in data_tr_xml, (
        "Data row missing <w:cantSplit/> element in OOXML"
    )
    assert "<w:tblHeader" not in data_tr_xml, "Data row must not have tblHeader attribute"


def test_messy_table_edge_cases(tmp_path: Path) -> None:
    """Verify ragged rows padding, escaped pipes, empty cells, and numeric fidelity."""
    # Table containing:
    # 1. Escaped pipe \| in header
    # 2. Uneven columns (row 1 has 3 cells, row 2 has 2 cells, row 3 has empty cells)
    # 3. Numeric values with decimals, commas, negative signs, percentages
    messy_md = (
        "| Metric \\| Name | Value | Percentage |\n"
        "|:---|---:|---:|\n"
        "| Total Revenue | 1,234,567.89 | +12.5% |\n"
        "| Ragged Row | -450.00 |\n"
        "| | empty cell test | |\n"
    )
    result = OCRResult(
        file_path="messy_table.pdf",
        pages=[PageResult(page_num=1, markdown=messy_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "messy_table.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    assert len(reopened.tables) == 1
    table = reopened.tables[0]

    # 1. Consistent column count across all rows (max_cols = 3)
    assert len(table.columns) == 3
    for idx, row in enumerate(table.rows):
        assert len(row.cells) == 3, f"Row {idx} does not have consistent 3 columns"

    # 2. Escaped pipe handling in header
    header_cell_0 = table.rows[0].cells[0].text
    assert "|" in header_cell_0, "Escaped pipe must be unescaped to literal '|'"
    assert "Metric | Name" in header_cell_0

    # 3. Numeric formatting verbatim preservation (decimals, commas, negative, percentage)
    row_1_val = table.rows[1].cells[1].text
    assert row_1_val == "1,234,567.89", f"Expected '1,234,567.89', got '{row_1_val}'"
    row_1_pct = table.rows[1].cells[2].text
    assert row_1_pct == "+12.5%"

    row_2_val = table.rows[2].cells[1].text
    assert row_2_val == "-450.00"

    # 4. Ragged row padding (row 2 was missing 3rd column, must be padded with empty cell)
    assert table.rows[2].cells[2].text == "", "Ragged row missing cell must be padded with empty string"

    # 5. Empty cells handling (row 3 has empty cells)
    assert table.rows[3].cells[0].text == "", "First cell of row 3 must be empty"
    assert table.rows[3].cells[1].text == "empty cell test"
    assert table.rows[3].cells[2].text == "", "Third cell of row 3 must be empty"


# Headings, Text Formatting, & Lists

def test_headings_mapping(tmp_path: Path) -> None:
    """Verify #, ##, ### map to Heading 1, Heading 2, Heading 3."""
    md = "# Main Title\n\n## Sub Section\n\n### Detail Level"
    result = OCRResult(
        file_path="headings.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "headings.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    assert paragraphs[0].text == "Main Title"
    assert paragraphs[0].style is not None
    assert paragraphs[0].style.name == "Heading 1"

    assert paragraphs[1].text == "Sub Section"
    assert paragraphs[1].style is not None
    assert paragraphs[1].style.name == "Heading 2"

    assert paragraphs[2].text == "Detail Level"
    assert paragraphs[2].style is not None
    assert paragraphs[2].style.name == "Heading 3"


def test_bold_italic_and_lists(tmp_path: Path) -> None:
    """Verify bold/italic run formatting and bullet/numbered list styles."""
    md = (
        "Text with **bold words** and *italic words*.\n\n"
        "- Bullet item 1\n"
        "- Bullet item 2\n\n"
        "1. First numbered\n"
        "2. Second numbered\n"
    )
    result = OCRResult(
        file_path="formatting.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "formatting.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Paragraph 0: bold and italic runs
    p0 = paragraphs[0]
    runs = p0.runs
    bold_runs = [r for r in runs if r.bold]
    italic_runs = [r for r in runs if r.italic]
    assert any("bold words" in r.text for r in bold_runs)
    assert any("italic words" in r.text for r in italic_runs)

    # Paragraphs 1-2: Bullet list
    assert paragraphs[1].style is not None
    assert paragraphs[1].style.name == "List Bullet"
    assert "Bullet item 1" in paragraphs[1].text
    assert paragraphs[2].style is not None
    assert paragraphs[2].style.name == "List Bullet"
    assert "Bullet item 2" in paragraphs[2].text

    # Paragraphs 3-4: Numbered list
    assert paragraphs[3].style is not None
    assert paragraphs[3].style.name == "List Number"
    assert "First numbered" in paragraphs[3].text
    assert paragraphs[4].style is not None
    assert paragraphs[4].style.name == "List Number"
    assert "Second numbered" in paragraphs[4].text


# Inline Code vs. Fenced Code Blocks & Blockquotes

def test_inline_code_vs_fenced_code_block_distinction(tmp_path: Path) -> None:
    """Verify inline code stays inline in same paragraph while fenced code is indented block."""
    md = (
        "Run the `cli.main` command with `--format` flag.\n\n"
        "```python\n"
        "def compute_total(a, b):\n"
        "    return a + b\n"
        "```\n\n"
        "> Supervised by Region Admin.\n"
    )
    result = OCRResult(
        file_path="code_distinction.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "code_distinction.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # 1. Inline code paragraph:
    # Must NOT have left_indent set to 0.5" / 36pt (it is normal body text)
    p_inline = paragraphs[0]
    assert p_inline.paragraph_format.left_indent is None or p_inline.paragraph_format.left_indent == 0
    # Must contain runs in the SAME paragraph with Consolas font
    consolas_runs = [r for r in p_inline.runs if r.font.name == "Consolas"]
    assert len(consolas_runs) >= 2
    assert any("cli.main" in r.text for r in consolas_runs)
    assert any("--format" in r.text for r in consolas_runs)
    # The normal text runs must also be in the SAME paragraph
    assert "Run the " in p_inline.text
    assert " command with " in p_inline.text

    # 2. Fenced code block paragraph:
    # Must be a standalone block with left indentation
    p_fenced = paragraphs[1]
    assert p_fenced.paragraph_format.left_indent == Pt(36), (
        f"Fenced code must have 36pt left indent, got {p_fenced.paragraph_format.left_indent}"
    )
    assert "def compute_total(a, b):" in p_fenced.text
    assert p_fenced.runs[0].font.name == "Consolas"

    # 3. Blockquote paragraph:
    # Must be a standalone block with left indentation and monospace font
    p_quote = paragraphs[2]
    assert p_quote.paragraph_format.left_indent == Pt(36), (
        f"Blockquote must have 36pt left indent, got {p_quote.paragraph_format.left_indent}"
    )
    assert "Supervised by Region Admin." in p_quote.text
    assert p_quote.runs[0].font.name == "Consolas"


# LaTeX Formula Preservation

def test_latex_formula_raw_text_preservation(tmp_path: Path) -> None:
    """Verify LaTeX math syntax ($...$ and $$...$$) is preserved verbatim as plain text."""
    latex_md = (
        "Inline formula: $E = mc^2$ and $x_{1,2} = \\frac{-b \\pm \\sqrt{b^2 - 4ac}}{2a}$.\n\n"
        "Block equation: $$\\sum_{i=1}^n x_i = X$$\n"
    )
    result = OCRResult(
        file_path="latex.pdf",
        pages=[PageResult(page_num=1, markdown=latex_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "latex.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Verify inline formulas remain verbatim strings with dollar signs
    p0_text = paragraphs[0].text
    assert "$E = mc^2$" in p0_text
    assert "$x_{1,2} = \\frac{-b \\pm \\sqrt{b^2 - 4ac}}{2a}$" in p0_text

    # Verify block formula remains verbatim string
    p1_text = paragraphs[1].text
    assert "$$\\sum_{i=1}^n x_i = X$$" in p1_text


# Export Helper (export_to_docx_bytes)

def test_export_to_docx_bytes() -> None:
    """Verify export_to_docx_bytes produces valid binary buffer."""
    result = OCRResult(
        file_path="helpers.pdf",
        pages=[PageResult(page_num=1, markdown="# Test Header\nSample content.", status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    data = export_to_docx_bytes(result)
    assert isinstance(data, bytes)
    assert len(data) > 0

    # Verify bytes can be loaded by Document
    doc_from_bytes = Document(io.BytesIO(data))
    assert doc_from_bytes.paragraphs[0].text == "Test Header"


def test_empty_or_failed_result_handled_cleanly() -> None:
    """Verify empty or failed OCRResult returns clean document without raising exceptions."""
    result_empty = OCRResult(file_path="empty.pdf", pages=[], status=JobStatus.FAILED)
    doc_empty = build_docx(result_empty)
    assert hasattr(doc_empty, "save")
    assert hasattr(doc_empty, "paragraphs")

    result_error = OCRResult(file_path="err.pdf", pages=[], status=JobStatus.FAILED, error="Corrupt PDF header")
    doc_err = build_docx(result_error)
    assert "Error: Corrupt PDF header" in doc_err.paragraphs[0].text


# Table Alignment, Inline Styling, and Degenerate Cases Tests

def test_table_cell_alignment_left_center_right(tmp_path: Path) -> None:
    """Verify table cell paragraph alignment matches GFM column specification (:---, :---:, ---:)."""
    table_md = (
        "| Left Header | Center Header | Right Header |\n"
        "|:---|:---:|---:|\n"
        "| L Data | C Data | R Data |\n"
    )
    result = OCRResult(
        file_path="align_table.pdf",
        pages=[PageResult(page_num=1, markdown=table_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "align_table.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    assert len(reopened.tables) == 1
    table = reopened.tables[0]

    # Row 0 (Headers)
    assert table.rows[0].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.rows[0].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.rows[0].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    # Row 1 (Data)
    assert table.rows[1].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.rows[1].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_table_cell_rich_inline_formatting(tmp_path: Path) -> None:
    """Verify bold, italic, strikethrough, and inline code inside table cells."""
    table_md = (
        "| Feature Header |\n"
        "|:---|\n"
        "| **Bold** and *Italic* and ~~Strike~~ and `code_snippet` |\n"
    )
    result = OCRResult(
        file_path="styled_table.pdf",
        pages=[PageResult(page_num=1, markdown=table_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "styled_table.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    table = reopened.tables[0]

    # Row 1 (Styled cells)
    p_styled = table.rows[1].cells[0].paragraphs[0]
    runs = p_styled.runs

    bold_runs = [r for r in runs if r.bold]
    assert any("Bold" in r.text for r in bold_runs)

    italic_runs = [r for r in runs if r.italic]
    assert any("Italic" in r.text for r in italic_runs)

    strike_runs = [r for r in runs if r.font.strike]
    assert any("Strike" in r.text for r in strike_runs)

    code_runs = [r for r in runs if r.font.name == "Consolas"]
    assert any("code_snippet" in r.text for r in code_runs)


def test_paragraph_softbreak_and_hardbreak(tmp_path: Path) -> None:
    """Verify soft breaks render as space and hard breaks render as newline in paragraphs."""
    md = "First Soft Line\nSecond Soft Line\n\nFirst Hard Line  \nSecond Hard Line"
    result = OCRResult(
        file_path="breaks.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )

    doc = build_docx(result)
    target = tmp_path / "breaks.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Paragraph 0: softbreak -> "First Soft Line Second Soft Line"
    assert "First Soft Line Second Soft Line" in paragraphs[0].text

    # Paragraph 1: hardbreak -> contains newline "\n"
    assert "First Hard Line\nSecond Hard Line" in paragraphs[1].text



def test_table_empty_and_zero_cols_handling() -> None:
    """Verify empty or zero-column table nodes exit cleanly without exceptions."""
    doc = Document()

    # 1. Empty table node
    root_empty = SyntaxTreeNode([Token("table_open", "table", 1), Token("table_close", "table", -1)])
    _render_table_node(doc, root_empty.children[0])
    assert len(doc.tables) == 0

    # 2. Table with empty row (max_cols == 0)
    root_zero = SyntaxTreeNode([
        Token("table_open", "table", 1),
        Token("tbody_open", "tbody", 1),
        Token("tr_open", "tr", 1),
        Token("tr_close", "tr", -1),
        Token("tbody_close", "tbody", -1),
        Token("table_close", "table", -1),
    ])
    _render_table_node(doc, root_zero.children[0])
    assert len(doc.tables) == 0


def test_render_markdown_page_empty_and_fallback_node() -> None:
    """Verify blank markdown strings and unrecognized block AST nodes are handled cleanly."""
    doc = Document()

    # 1. Whitespace only markdown string
    render_markdown_page(doc, "   \n\t  \n")
    assert len(doc.paragraphs) == 0

    # 2. Unrecognized block node with content
    root_content = SyntaxTreeNode([Token("custom_unknown_block", "", 0, content="Fallback raw content string")])
    _render_ast_node(doc, root_content.children[0])
    assert doc.paragraphs[0].text == "Fallback raw content string"

    # 3. Unrecognized block node with children
    doc2 = Document()
    root_container = SyntaxTreeNode([
        Token("custom_container_open", "div", 1),
        Token("paragraph_open", "p", 1),
        Token("paragraph_close", "p", -1),
        Token("custom_container_close", "div", -1),
    ])
    _render_ast_node(doc2, root_container.children[0])
    assert len(doc2.paragraphs) == 1


def test_multipage_with_failed_or_empty_pages_page_break_fidelity(tmp_path: Path) -> None:
    """Verify page breaks are placed strictly between valid, non-empty pages."""
    result = OCRResult(
        file_path="multipage_mixed.pdf",
        pages=[
            PageResult(page_num=1, markdown="# Page 1 Valid\nContent.", status=JobStatus.SUCCESS),
            PageResult(page_num=2, markdown="", status=JobStatus.FAILED, error="Timeout"),
            PageResult(page_num=3, markdown="   \n  ", status=JobStatus.SUCCESS),
            PageResult(page_num=4, markdown="# Page 4 Valid\nFinal content.", status=JobStatus.SUCCESS),
        ],
        status=JobStatus.PARTIAL,
    )

    doc = build_docx(result)
    target = tmp_path / "multipage_mixed.docx"
    doc.save(str(target))

    reopened = Document(str(target))
    paragraphs = reopened.paragraphs

    # Count literal page breaks (<w:br w:type="page"/>)
    page_break_count = sum(
        1 for p in paragraphs for r in p.runs
        if '<w:br w:type="page"/>' in r._r.xml
    )

    # Exactly 2 valid pages with content (Page 1 and Page 4) -> exactly 1 page break between them
    assert page_break_count == 1, f"Expected exactly 1 page break, found {page_break_count}"
    assert "Page 1 Valid" in paragraphs[0].text
    # Trailing paragraph must not end with a page break
    assert '<w:br w:type="page"/>' not in paragraphs[-1].runs[-1]._r.xml


# Nested Lists & Image Placeholder Rendering

def test_no_text_silently_dropped() -> None:
    """Every text node in the markdown AST must appear somewhere in the DOCX output.

    Covers nested bullet lists (3 levels), nested ordered lists (2 levels),
    and image references with alt text.
    """
    from markdown_it import MarkdownIt

    md = (
        "- level1\n"
        "  - level2\n"
        "    - level3\n"
        "- another1\n"
        "\n"
        "1. first ordered\n"
        "   1. nested ordered\n"
        "2. second ordered\n"
        "\n"
        "![my alt text](img.png)\n"
    )

    # Collect every non-empty text node from the AST
    def collect_ast_texts(node: SyntaxTreeNode) -> list[str]:
        texts = []
        if node.type == "text" and node.content:
            texts.append(node.content)
        for child in node.children:
            texts.extend(collect_ast_texts(child))
        return texts

    parser = MarkdownIt()
    tokens = parser.parse(md)
    root = SyntaxTreeNode(tokens)
    ast_texts: list[str] = []
    for child in root.children:
        ast_texts.extend(collect_ast_texts(child))

    result = OCRResult(
        file_path="nested_lists.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)

    # Concatenate all paragraph text from the DOCX
    docx_text = " ".join(p.text for p in doc.paragraphs)

    for text in ast_texts:
        assert text in docx_text, (
            f"AST text node {text!r} is missing from DOCX output. "
            f"Full DOCX text: {docx_text!r}"
        )


def test_image_node_renders_as_placeholder() -> None:
    """Image references must produce a visible [image: <alt>] placeholder, not bare alt text."""
    md = "Before image. ![diagram of process](flow.png) After image."
    result = OCRResult(
        file_path="img_test.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    para_text = doc.paragraphs[0].text

    # Placeholder must appear with bracket formatting
    assert "[image: diagram of process]" in para_text, (
        f"Expected '[image: diagram of process]' placeholder in paragraph text, got: {para_text!r}"
    )
    # The bare alt text must NOT appear as plain prose outside the brackets
    stripped = para_text.replace("[image: diagram of process]", "")
    assert "diagram of process" not in stripped, (
        f"Alt text leaked as bare prose outside the placeholder: {para_text!r}"
    )


# HTML Table Export Tests

def test_html_table_synthetic_fixtures(tmp_path: Path) -> None:
    """Verify receipt fixture and wide table fixture render as single Word table with matching rows, cols, and text."""
    receipt_html = (
        "<table>\n"
        "<thead>\n"
        '<tr><th>Item</th><th style="text-align: right">Price</th></tr>\n'
        "</thead>\n"
        "<tbody>\n"
        '<tr><td>Espresso</td><td style="text-align: right">$3.50</td></tr>\n'
        '<tr><td>Croissant</td><td style="text-align: right">$4.00</td></tr>\n'
        "</tbody>\n"
        "</table>"
    )
    result_receipt = OCRResult(
        file_path="receipt.pdf",
        pages=[PageResult(page_num=1, markdown=receipt_html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc_receipt = build_docx(result_receipt)
    assert len(doc_receipt.tables) == 1
    t_receipt = doc_receipt.tables[0]
    assert len(t_receipt.rows) == 3
    assert len(t_receipt.columns) == 2
    assert t_receipt.rows[0].cells[0].text == "Item"
    assert t_receipt.rows[0].cells[1].text == "Price"
    assert t_receipt.rows[1].cells[0].text == "Espresso"
    assert t_receipt.rows[1].cells[1].text == "$3.50"
    assert t_receipt.rows[2].cells[0].text == "Croissant"
    assert t_receipt.rows[2].cells[1].text == "$4.00"

    wide_html = (
        '<table class="table table-bordered">\n'
        "<thead>\n"
        "<tr><th>ID</th><th>First Name</th><th>Last Name</th><th>Role</th><th>Status</th></tr>\n"
        "</thead>\n"
        "<tbody>\n"
        "<tr><td>001</td><td>Alice</td><td>Smith</td><td>Admin</td><td>Active</td></tr>\n"
        "<tr><td>002</td><td>Bob</td><td>Jones</td><td>User</td><td>Inactive</td></tr>\n"
        "</tbody>\n"
        "</table>"
    )
    result_wide = OCRResult(
        file_path="wide.pdf",
        pages=[PageResult(page_num=1, markdown=wide_html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc_wide = build_docx(result_wide)
    assert len(doc_wide.tables) == 1
    t_wide = doc_wide.tables[0]
    assert len(t_wide.rows) == 3
    assert len(t_wide.columns) == 5
    assert [c.text for c in t_wide.rows[0].cells] == ["ID", "First Name", "Last Name", "Role", "Status"]
    assert [c.text for c in t_wide.rows[1].cells] == ["001", "Alice", "Smith", "Admin", "Active"]
    assert [c.text for c in t_wide.rows[2].cells] == ["002", "Bob", "Jones", "User", "Inactive"]


def test_html_table_with_blank_lines() -> None:
    """Verify HTML table with empty lines between rows renders as single Word table with all rows preserved."""
    md = (
        "<table>\n\n"
        "<thead>\n"
        "<tr><th>Region</th><th>Target</th></tr>\n"
        "</thead>\n\n"
        "<tbody>\n\n"
        "<tr><td>North</td><td>100</td></tr>\n\n"
        "<tr><td>South</td><td>200</td></tr>\n\n"
        "</tbody>\n"
        "</table>"
    )
    result = OCRResult(
        file_path="blank_lines.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.rows[0].cells[0].text == "Region"
    assert table.rows[1].cells[0].text == "North"
    assert table.rows[2].cells[0].text == "South"


def test_html_table_truncated_unclosed() -> None:
    """Verify truncated HTML table ending mid-cell/row does not crash and preserves parsed rows."""
    truncated_md = (
        "<table>\n"
        "<tr><th>Header 1</th><th>Header 2</th></tr>\n"
        "<tr><td>Row 1 Cell 1</td><td>Row 1 Cell 2</td></tr>\n"
        "<tr><td>Truncated Cell"
    )
    result = OCRResult(
        file_path="truncated.pdf",
        pages=[PageResult(page_num=1, markdown=truncated_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "Header 1"
    assert table.rows[1].cells[0].text == "Row 1 Cell 1"
    assert table.rows[2].cells[0].text == "Truncated Cell"


def test_html_table_entities_and_merged_cells() -> None:
    """Verify character entities (&amp;, &gt;) unescape and colspan/rowspan apply Word table merges."""
    table_md = (
        "<table>\n"
        '<tr><th colspan="2">Q1 &amp; Q2 Summary</th><th>Notes &gt; Remarks</th></tr>\n'
        '<tr><td rowspan="2">Group A</td><td>Metric 1</td><td>Pass</td></tr>\n'
        '<tr><td>Metric 2</td><td>Pass</td></tr>\n'
        "</table>"
    )
    result = OCRResult(
        file_path="merged.pdf",
        pages=[PageResult(page_num=1, markdown=table_md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 3

    # Unescaped entities & colspan merge
    assert table.rows[0].cells[0].text == "Q1 & Q2 Summary"
    assert table.rows[0].cells[1].text == "Q1 & Q2 Summary"
    assert table.rows[0].cells[2].text == "Notes > Remarks"

    # Rowspan merge
    assert table.rows[1].cells[0].text == "Group A"
    assert table.rows[2].cells[0].text == "Group A"
    assert table.rows[1].cells[1].text == "Metric 1"
    assert table.rows[2].cells[1].text == "Metric 2"


def test_html_block_non_table_text_preservation() -> None:
    """Verify non-table text around HTML tables is preserved as paragraphs alongside the table."""
    md = (
        "<table>\n"
        "<tr><td>Row 1</td></tr>\n"
        "</table>\n"
        "Footnote: text after table"
    )
    result = OCRResult(
        file_path="footnote.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    assert any("Footnote: text after table" in p.text for p in doc.paragraphs)


def test_html_table_hostile_spans() -> None:
    """Verify hostile/malformed colspan and rowspan do not cause memory exhaustion or crash, and columns are clamped to <= 63."""
    hostile_html = (
        "<table\n"
        '<tr><th colspan="100000" rowspan="-5">Header</th><th colspan="invalid">Col2</th></tr>\n'
        '<tr><td rowspan="99999">Cell1</td><td>Cell2</td></tr>\n'
        "</table>"
    )
    result = OCRResult(
        file_path="hostile.pdf",
        pages=[PageResult(page_num=1, markdown=hostile_html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.columns) <= 63
    assert table.rows[0].cells[0].text == "Header"


def test_html_table_truncated_with_blank_lines_and_trailing_markdown() -> None:
    """Verify truncated HTML table with blank lines and trailing markdown does not swallow trailing prose."""
    md = (
        "<table>\n"
        "<tr><th>Col 1</th><th>Col 2</th></tr>\n\n"
        "<tr><td>Row 1</td><td>Data 1</td></tr>\n\n"
        "This is normal trailing prose that should remain a paragraph."
    )
    result = OCRResult(
        file_path="truncated_trailing.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2
    assert any(
        p.text == "This is normal trailing prose that should remain a paragraph."
        for p in doc.paragraphs
    )


def test_html_table_real_captured_fixtures() -> None:
    """Verify real captured receipt and wide table fixtures render exact rows, cols, entities, and values."""
    receipt_html = (
        "<table>"
        "<thead><tr><th>Item Description</th><th>Qty</th><th>Rate</th><th>Subtotal</th></tr></thead>"
        "<tbody>"
        "<tr><td>Thermal Transfer Ribbon 110mm x 300m (Pack of 2)</td><td>2</td><td>$12.00</td><td>$24.00</td></tr>"
        "<tr><td>Precision Optical Sensor Calibration Pack</td><td>5</td><td>$25.00</td><td>$125.00</td></tr>"
        "<tr><td>USB-C 3.2 Gen2 High-Speed Host Interface Cable 2m</td><td>1</td><td>$18.50</td><td>$18.50</td></tr>"
        "</tbody>"
        "<tfoot>"
        "<tr><td>Subtotal:</td><td></td><td>$167.50</td><td></td></tr>"
        "<tr><td>Goods &amp; Services Tax (GST 9%)</td><td></td><td>$15.08</td><td></td></tr>"
        "<tr><td>Total Amount Paid:</td><td></td><td>$182.58</td><td></td></tr>"
        "</tfoot>"
        "</table>"
    )
    result_receipt = OCRResult(
        file_path="real_receipt.pdf",
        pages=[PageResult(page_num=1, markdown=receipt_html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc_receipt = build_docx(result_receipt)
    assert len(doc_receipt.tables) == 1
    t_receipt = doc_receipt.tables[0]
    assert len(t_receipt.rows) == 7
    assert len(t_receipt.columns) == 4
    assert t_receipt.rows[5].cells[0].text == "Goods & Services Tax (GST 9%)"
    assert t_receipt.rows[6].cells[2].text == "$182.58"

    wide_html = (
        '<table class="table table-bordered">'
        "<thead>"
        "<tr><th>Region / Unit</th><th>Gross Billings</th><th>Direct Costs</th><th>Gross Margin</th><th>Overhead</th><th>Net Operating Profit</th><th>Margin %</th></tr>"
        "</thead>"
        "<tbody>"
        "<tr><td>North America East</td><td>$4,850,000</td><td>$2,100,000</td><td>$2,750,000</td><td>$890,000</td><td>$1,860,000</td><td>38.4%</td></tr>"
        "<tr><td>North America West</td><td>$5,240,000</td><td>$2,350,000</td><td>$2,890,000</td><td>$940,000</td><td>$1,950,000</td><td>37.2%</td></tr>"
        "<tr><td>Western Europe (EMEA)</td><td>$3,920,000</td><td>$1,880,000</td><td>$2,040,000</td><td>$710,000</td><td>$1,330,000</td><td>33.9%</td></tr>"
        "<tr><td>Asia-Pacific (APAC)</td><td>$4,150,000</td><td>$1,620,000</td><td>$2,530,000</td><td>$680,000</td><td>$1,850,000</td><td>44.6%</td></tr>"
        "<tr><td>Latin America (LATAM)</td><td>$1,450,000</td><td>$820,000</td><td>$630,000</td><td>$340,000</td><td>$290,000</td><td>20.0%</td></tr>"
        "<tr><td>Consolidated Total</td><td>$19,610,000</td><td>$8,770,000</td><td>$10,840,000</td><td>$3,560,000</td><td>$7,280,000</td><td>37.1%</td></tr>"
        "</tbody>"
        "</table>"
    )
    result_wide = OCRResult(
        file_path="real_wide.pdf",
        pages=[PageResult(page_num=1, markdown=wide_html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc_wide = build_docx(result_wide)
    assert len(doc_wide.tables) == 1
    t_wide = doc_wide.tables[0]
    assert len(t_wide.rows) == 7
    assert len(t_wide.columns) == 7
    assert t_wide.rows[0].cells[0].text == "Region / Unit"
    assert t_wide.rows[6].cells[1].text == "$19,610,000"


def test_html_table_cell_paragraphs() -> None:
    """Verify <p> tags inside table cells are separated by newlines."""
    md = (
        "<table>\n"
        "<tr><td><p>First paragraph</p><p>Second paragraph</p></td></tr>\n"
        "</table>"
    )
    result = OCRResult(
        file_path="cell_p.pdf",
        pages=[PageResult(page_num=1, markdown=md, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    cell_text = doc.tables[0].rows[0].cells[0].text
    assert cell_text == "First paragraph\nSecond paragraph"


def test_html_table_docx_roundtrip_disk_reopen_and_ooxml_structure(tmp_path: Path) -> None:
    """Verify clamped table DOCX disk round-trip and inspect raw OOXML elements."""
    html = (
        "<table>\n"
        '<tr><td colspan="100" rowspan="5">Clamped Spans Cell</td><td>Col 2</td></tr>\n'
        "<tr><td>Row 2</td></tr>\n"
        "<tr><td>Row 3</td></tr>\n"
        "<tr><td>Row 4</td></tr>\n"
        "<tr><td>Row 5</td></tr>\n"
        "</table>"
    )
    result = OCRResult(
        file_path="clamped.pdf",
        pages=[PageResult(page_num=1, markdown=html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    file_path = tmp_path / "clamped.docx"
    doc.save(str(file_path))

    # Re-open with docx.Document
    reopened = docx.Document(str(file_path))
    assert len(reopened.tables) == 1
    table = reopened.tables[0]
    assert len(table.columns) <= 63
    assert table.rows[0].cells[0].text == "Clamped Spans Cell"

    # Unzip clamped.docx and parse word/document.xml
    with zipfile.ZipFile(str(file_path)) as z:
        xml_content = z.read("word/document.xml")
    root = ET.fromstring(xml_content)

    # Asserts XML root tag
    assert root.tag.endswith("document")

    # Asserts w:tbl and w:tblGrid exist
    tbls = [el for el in root.iter() if el.tag.endswith("tbl")]
    assert len(tbls) >= 1
    tbl_grids = [el for el in root.iter() if el.tag.endswith("tblGrid")]
    assert len(tbl_grids) >= 1

    # Asserts w:gridCol count <= 63
    grid_cols = [el for el in root.iter() if el.tag.endswith("gridCol")]
    assert 0 < len(grid_cols) <= 63

    # Asserts w:hMerge (w:gridSpan) and w:vMerge OOXML tags exist on merged cells
    has_h_merge = any(
        el.tag.endswith(("gridSpan", "hMerge")) or "gridSpan" in el.tag or "hMerge" in el.tag
        for el in root.iter()
    )
    has_v_merge = any(
        el.tag.endswith("vMerge") or "vMerge" in el.tag
        for el in root.iter()
    )
    assert has_h_merge, "Expected horizontal merge tag (w:gridSpan or w:hMerge) in OOXML"
    assert has_v_merge, "Expected vertical merge tag (w:vMerge) in OOXML"


def test_html_table_row_count_clamping() -> None:
    """Verify HTML tables exceeding max_rows are clamped with a visible DOCX notice."""
    rows_html = "".join(f"<tr><td>Row {i}</td></tr>\n" for i in range(25))
    html = f"<table>\n{rows_html}</table>"
    doc = Document()
    render_markdown_page(doc, html, max_html_table_rows=5)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 5
    assert any("[Table truncated: 20 rows omitted to fit document limits]" in p.text for p in doc.paragraphs)

    large_rows = "".join(f"<tr><td>Row {i}</td></tr>\n" for i in range(1005))
    extractor = _HTMLTableExtractor()
    extractor.feed(f"<table>\n{large_rows}</table>")
    assert len(extractor.tables[0]) == 1000
    assert extractor.sequence[0][2] == 5


def test_html_table_nested_table_in_cell_safety() -> None:
    """Verify parser safely flattens nested tables without crash and preserves cell text."""
    html = "<table><tr><td>Outer Cell <table><tr><td>Inner Cell</td></tr></table> More Outer</td></tr></table>"
    result = OCRResult(
        file_path="nested_table.pdf",
        pages=[PageResult(page_num=1, markdown=html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 1
    assert len(doc.tables[0].columns) == 1
    assert doc.tables[0].rows[0].cells[0].text == "Outer Cell Inner Cell More Outer"


def test_html_table_column_clamping_visible_notice() -> None:
    """Verify HTML table exceeding 63 columns clamps to 63 and emits visible truncation notice."""
    cells = "".join(f"<td>C{i}</td>" for i in range(65))
    html = f"<table><tr>{cells}</tr></table>"
    doc = Document()
    render_markdown_page(doc, html)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 63
    assert any("[Table truncated: 2 columns omitted to fit document limits]" in p.text for p in doc.paragraphs)


def test_gfm_table_column_clamping_visible_notice() -> None:
    """Verify GFM pipe table with 65 columns clamps to 63 and emits visible truncation notice."""
    header = "| " + " | ".join(f"Col{i}" for i in range(65)) + " |"
    delimiter = "| " + " | ".join("---" for _ in range(65)) + " |"
    row = "| " + " | ".join(f"Val{i}" for i in range(65)) + " |"
    table_md = f"{header}\n{delimiter}\n{row}\n"
    doc = Document()
    render_markdown_page(doc, table_md)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 63
    assert any("[Table truncated: 2 columns omitted to fit document limits]" in p.text for p in doc.paragraphs)


def test_html_table_colspan_crossing_column_limit() -> None:
    """Verify cell at col 60 with colspan 6 (total 66) clamps grid to 63 and reports 3 omitted cols."""
    first_60 = "".join(f"<td>C{i}</td>" for i in range(60))
    wide_cell = '<td colspan="6">Wide</td>'
    html = f"<table><tr>{first_60}{wide_cell}</tr></table>"
    doc = Document()
    render_markdown_page(doc, html)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 63
    assert any("[Table truncated: 3 columns omitted to fit document limits]" in p.text for p in doc.paragraphs)


def test_row_cap_boundary_1000_vs_1001() -> None:
    """Verify boundary at exactly 1000 rows (0 omitted) vs 1001 rows (1 row omitted with notice)."""
    # 1000 rows boundary
    ext_1000 = _HTMLTableExtractor(max_rows=1000)
    ext_1000.feed("<table>" + "".join(f"<tr><td>R{i}</td></tr>" for i in range(1000)) + "</table>")
    ext_1000.close()
    assert len(ext_1000.tables[0]) == 1000
    assert ext_1000.sequence[0][2] == 0

    doc_1000 = Document()
    _render_parsed_html_table(doc_1000, [[{"text": f"R{i}"}] for i in range(5)], truncated_rows=ext_1000.sequence[0][2])
    assert not any("[Table truncated" in p.text for p in doc_1000.paragraphs)

    # 1001 rows boundary
    ext_1001 = _HTMLTableExtractor(max_rows=1000)
    ext_1001.feed("<table>" + "".join(f"<tr><td>R{i}</td></tr>" for i in range(1001)) + "</table>")
    ext_1001.close()
    assert len(ext_1001.tables[0]) == 1000
    assert ext_1001.sequence[0][2] == 1

    doc_1001 = Document()
    _render_parsed_html_table(doc_1001, [[{"text": f"R{i}"}] for i in range(5)], truncated_rows=ext_1001.sequence[0][2])
    assert any("[Table truncated: 1 rows omitted to fit document limits]" in p.text for p in doc_1001.paragraphs)


def test_html_table_nested_table_text_separator() -> None:
    """Verify nested table with multiple cells and rows separates cell texts with spaces and rows with newlines."""
    html = (
        "<table><tr><td>"
        "Outer "
        "<table>"
        "<tr><td>Inner1</td><td>Inner2</td></tr>"
        "<tr><td>Inner3</td><td>Inner4</td></tr>"
        "</table>"
        "</td></tr></table>"
    )
    result = OCRResult(
        file_path="nested_multi.pdf",
        pages=[PageResult(page_num=1, markdown=html, status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
    )
    doc = build_docx(result)
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "Outer Inner1 Inner2\nInner3 Inner4"


def test_table_both_rows_and_cols_truncated_notice() -> None:
    """Verify notice phrasing when both rows and columns exceed document limits."""
    doc = Document()
    _append_table_truncation_note(doc, truncated_rows=5, truncated_cols=2)
    assert any("[Table truncated: 5 rows and 2 columns omitted to fit document limits]" in p.text for p in doc.paragraphs)



