"""Unit tests for output formatting and artifact saving (core/formatter.py)."""

import io
import json
from pathlib import Path
from unittest.mock import patch
import pytest
from docx import Document

from core.formatter import (
    WINDOWS_RESERVED_NAMES,
    format_output,
    resolve_unique_stem,
    sanitize_export_path,
    sanitize_filename_stem,
    save_artifacts,
)
from core.models import JobConfig, JobStatus, OCRResult, OutputFormat, PageResult


# ==============================================================================
# Fixtures
# ==============================================================================

@pytest.fixture
def sample_ocr_result() -> OCRResult:
    result = OCRResult(
        file_path="reports/financial_audit.pdf",
        pages=[
            PageResult(
                page_num=1,
                markdown="# Financial Report 2026\nQuarterly Revenue: $1.2M",
                raw_json={"id": "cmpl-1"},
                latency=0.3,
                status=JobStatus.SUCCESS,
            ),
            PageResult(
                page_num=2,
                markdown="## Expenses Breakdown\nOperational Cost: $400K",
                raw_json={"id": "cmpl-2"},
                latency=0.25,
                status=JobStatus.SUCCESS,
            ),
        ],
        total_duration=0.55,
        status=JobStatus.SUCCESS,
    )
    return result


# ==============================================================================
# Format Output Tests
# ==============================================================================

def test_format_output_markdown_only(sample_ocr_result: OCRResult) -> None:
    outputs = format_output(sample_ocr_result, OutputFormat.MARKDOWN)

    assert "markdown" in outputs
    assert "json" not in outputs
    assert "# Financial Report 2026" in outputs["markdown"]
    assert "## Expenses Breakdown" in outputs["markdown"]
    assert "---" in outputs["markdown"]


def test_format_output_json_only(sample_ocr_result: OCRResult) -> None:
    outputs = format_output(sample_ocr_result, OutputFormat.JSON)

    assert "json" in outputs
    assert "markdown" not in outputs

    parsed = json.loads(outputs["json"])
    assert parsed["file_path"] == "reports/financial_audit.pdf"
    assert parsed["status"] == "SUCCESS"
    assert parsed["page_count"] == 2


def test_format_output_both(sample_ocr_result: OCRResult) -> None:
    outputs = format_output(sample_ocr_result, OutputFormat.BOTH)

    assert "markdown" in outputs
    assert "json" in outputs
    assert "# Financial Report 2026" in outputs["markdown"]
    assert json.loads(outputs["json"])["page_count"] == 2


# ==============================================================================
# Save Artifacts Tests
# ==============================================================================

def test_save_artifacts_default_stem_from_file_path(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    config = JobConfig(output_format=OutputFormat.BOTH)
    saved = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)

    expected_md = tmp_path / "financial_audit.md"
    expected_json = tmp_path / "financial_audit.json"

    assert saved["markdown"] == expected_md.resolve()
    assert saved["json"] == expected_json.resolve()

    assert expected_md.exists()
    assert expected_json.exists()

    md_content = expected_md.read_text(encoding="utf-8")
    assert "# Financial Report 2026" in md_content

    json_data = json.loads(expected_json.read_text(encoding="utf-8"))
    assert json_data["status"] == "SUCCESS"


def test_save_artifacts_custom_base_name_and_format_markdown_only(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    config = JobConfig(output_format=OutputFormat.MARKDOWN)
    saved = save_artifacts(
        sample_ocr_result,
        config,
        output_dir=tmp_path,
        base_name="custom_export_name",
    )

    assert "markdown" in saved
    assert "json" not in saved

    expected_md = tmp_path / "custom_export_name.md"
    assert saved["markdown"] == expected_md.resolve()
    assert expected_md.exists()
    assert not (tmp_path / "custom_export_name.json").exists()


def test_save_artifacts_in_memory_fallback_stem(tmp_path: Path) -> None:
    result = OCRResult(
        file_path="<in-memory>",
        pages=[PageResult(page_num=1, markdown="In-memory page", status=JobStatus.SUCCESS)],
    )
    config = JobConfig(output_format=OutputFormat.JSON)

    saved = save_artifacts(result, config, output_dir=tmp_path)
    expected_json = tmp_path / "ocr_result.json"

    assert saved["json"] == expected_json.resolve()
    assert expected_json.exists()


# ==============================================================================
# Filename Stem Sanitization Tests (Finding 1.2)
# ==============================================================================

@pytest.mark.parametrize(
    ("raw_stem", "expected"),
    [
        ("report:january*2026?", "report_january_2026_"),
        ("path/to\\bad|name<test>", "path_to_bad_name_test_"),
        ("trailing_dots....", "trailing_dots"),
        ("trailing_spaces   ", "trailing_spaces"),
        ("trailing_dots_and_spaces. . . ", "trailing_dots_and_spaces"),
        ("   ", "ocr_result"),
        ("....", "ocr_result"),
        ("CON", "_CON"),
        ("con", "_con"),
        ("PRN", "_PRN"),
        ("aux", "_aux"),
        ("NUL", "_NUL"),
        ("com1", "_com1"),
        ("COM9", "_COM9"),
        ("lpt1", "_lpt1"),
        ("LPT9", "_LPT9"),
    ],
)
def test_sanitize_filename_stem(raw_stem: str, expected: str) -> None:
    assert sanitize_filename_stem(raw_stem) == expected


# ==============================================================================
# Unique Collision Resolution Tests (Finding 1.1)
# ==============================================================================

def test_resolve_unique_stem_with_existing_disk_files(tmp_path: Path) -> None:
    # Pre-create disk files
    (tmp_path / "report.md").write_text("existing", encoding="utf-8")
    (tmp_path / "report_2.json").write_text("{}", encoding="utf-8")

    stem = resolve_unique_stem("report", output_dir=tmp_path)
    assert stem == "report_3"


def test_resolve_unique_stem_with_batch_used_stems(tmp_path: Path) -> None:
    used_stems = set()
    stem1 = resolve_unique_stem("invoice", output_dir=tmp_path, used_stems=used_stems)
    assert stem1 == "invoice"
    assert "invoice" in used_stems

    stem2 = resolve_unique_stem("invoice", output_dir=tmp_path, used_stems=used_stems)
    assert stem2 == "invoice_2"
    assert "invoice_2" in used_stems

    stem3 = resolve_unique_stem("invoice", output_dir=tmp_path, used_stems=used_stems)
    assert stem3 == "invoice_3"
    assert "invoice_3" in used_stems


def test_save_artifacts_avoids_clobbering_existing_disk_files(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    config = JobConfig(output_format=OutputFormat.MARKDOWN)

    # First save
    saved1 = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)
    assert saved1["markdown"] == (tmp_path / "financial_audit.md").resolve()

    # Second save with same result (without used_stems) should still avoid overwriting
    saved2 = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)
    assert saved2["markdown"] == (tmp_path / "financial_audit_2.md").resolve()

    assert (tmp_path / "financial_audit.md").exists()
    assert (tmp_path / "financial_audit_2.md").exists()


# ==============================================================================
# Path Privacy Export Sanitization Tests (Finding 4.1)
# ==============================================================================

def test_sanitize_export_path_relative_to_base_and_cwd(tmp_path: Path) -> None:
    """Verify sanitize_export_path relativizes within base_dir or cwd."""
    sub_file = tmp_path / "batches" / "scan1.pdf"

    # Relativize against explicit base_dir
    sanitized = sanitize_export_path(sub_file, base_dir=tmp_path)
    assert sanitized == "batches/scan1.pdf"

    # In-memory marker preserved
    assert sanitize_export_path("<in-memory>") == "<in-memory>"
    assert sanitize_export_path("<bytes>") == "<bytes>"
    assert sanitize_export_path("") == ""


def test_save_artifacts_sanitizes_path_default(tmp_path: Path) -> None:
    """Verify save_artifacts sanitizes file_path and error strings in exported JSON by default without mutating OCRResult."""
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    error_msg = f"File is empty (0 bytes): '{raw_path}'"
    page_error_msg = f"Failed to render frame 1 for '{raw_path}'"
    result = OCRResult(
        file_path=str(raw_path),
        pages=[PageResult(page_num=1, markdown="Secret", status=JobStatus.FAILED, error=page_error_msg)],
        error=error_msg,
        status=JobStatus.FAILED,
    )
    config = JobConfig(output_format=OutputFormat.JSON)

    # 1. Default (sanitize_path=True) with custom base_dir
    saved = save_artifacts(
        result,
        config,
        output_dir=tmp_path / "out",
        base_dir=tmp_path / "private_user_dir",
    )
    exported_json = json.loads(saved["json"].read_text(encoding="utf-8"))
    assert exported_json["file_path"] == "secret_doc.pdf"
    assert exported_json["error"] == "File is empty (0 bytes): 'secret_doc.pdf'"
    assert exported_json["pages"][0]["error"] == "Failed to render frame 1 for 'secret_doc.pdf'"
    assert str(raw_path) not in exported_json["error"]
    assert str(raw_path) not in exported_json["pages"][0]["error"]

    # Verify OCRResult internal data model was NOT mutated
    assert result.file_path == str(raw_path)
    assert result.error == error_msg
    assert result.to_dict()["file_path"] == str(raw_path)
    assert result.to_dict()["error"] == error_msg

    # 2. Explicit sanitize_path=False preserves absolute path and error
    saved_raw = save_artifacts(
        result,
        config,
        output_dir=tmp_path / "out_raw",
        sanitize_path=False,
    )
    exported_raw_json = json.loads(saved_raw["json"].read_text(encoding="utf-8"))
    assert exported_raw_json["file_path"] == str(raw_path)
    assert exported_raw_json["error"] == error_msg
    assert exported_raw_json["pages"][0]["error"] == page_error_msg


def test_save_artifacts_atomic_write_preserves_existing_on_error(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts writes atomically and unlinks .tmp without corrupting existing file on error (C-6)."""
    config = JobConfig(output_format=OutputFormat.BOTH)

    # 1. Clean run writes files atomically
    saved = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)
    md_file = saved["markdown"]
    json_file = saved["json"]
    assert md_file.exists()
    assert json_file.exists()
    assert not list(tmp_path.glob("*.tmp"))

    orig_md_content = md_file.read_text(encoding="utf-8")

    # 2. Simulate write_text failure on next save
    real_write_text = Path.write_text

    def failing_write_text(self, data, *args, **kwargs):
        if str(self).endswith(".tmp"):
            raise OSError("Simulated disk error during temp write")
        return real_write_text(self, data, *args, **kwargs)

    with patch.object(Path, "write_text", side_effect=failing_write_text, autospec=True):
        with pytest.raises(OSError, match="Simulated disk error"):
            save_artifacts(sample_ocr_result, config, output_dir=tmp_path)

    # Verify original file was not corrupted and no lingering .tmp files remain
    assert md_file.read_text(encoding="utf-8") == orig_md_content
    assert not list(tmp_path.glob("*.tmp"))


def test_save_artifacts_json_atomic_replace_preserves_existing_on_error(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts unlinks .json.tmp and preserves existing .json on replace failure."""
    config = JobConfig(output_format=OutputFormat.JSON)

    saved = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)
    json_file = saved["json"]
    assert json_file.exists()
    orig_json_content = json_file.read_text(encoding="utf-8")

    real_replace = Path.replace

    def failing_replace(self, target, *args, **kwargs):
        if str(self).endswith(".json.tmp"):
            raise OSError("Simulated disk error during json replace")
        return real_replace(self, target, *args, **kwargs)

    with patch.object(Path, "replace", side_effect=failing_replace, autospec=True):
        with pytest.raises(OSError, match="Simulated disk error during json replace"):
            save_artifacts(sample_ocr_result, config, output_dir=tmp_path)

    # Verify original json file was not corrupted and no lingering .tmp files remain
    assert json_file.read_text(encoding="utf-8") == orig_json_content
    assert not list(tmp_path.glob("*.tmp"))


def test_format_output_docx_returns_raw_bytes(sample_ocr_result: OCRResult) -> None:
    """Verify format_output returns raw bytes for OutputFormat.DOCX without str encoding."""
    outputs = format_output(sample_ocr_result, OutputFormat.DOCX)

    assert "docx" in outputs
    assert "markdown" not in outputs
    assert "json" not in outputs
    assert isinstance(outputs["docx"], bytes)
    assert len(outputs["docx"]) > 0

    # Verify binary payload is a valid DOCX loadable by python-docx Document()
    doc = Document(io.BytesIO(outputs["docx"]))
    assert len(doc.paragraphs) > 0


def test_format_output_both_untouched_never_produces_docx(sample_ocr_result: OCRResult) -> None:
    """Verify OutputFormat.BOTH still strictly produces markdown and json only, never docx."""
    outputs = format_output(sample_ocr_result, OutputFormat.BOTH)

    assert "markdown" in outputs
    assert "json" in outputs
    assert "docx" not in outputs


def test_resolve_unique_stem_checks_existing_docx(tmp_path: Path) -> None:
    """Verify resolve_unique_stem increments suffix when an existing .docx file exists on disk."""
    # Pre-create foo.docx
    existing_docx = tmp_path / "foo.docx"
    existing_docx.write_bytes(b"mock docx binary content")

    stem = resolve_unique_stem("foo", output_dir=tmp_path)
    assert stem == "foo_2", f"Expected 'foo_2' when 'foo.docx' exists, got '{stem}'"

    # Pre-create foo_2.docx as well
    (tmp_path / "foo_2.docx").write_bytes(b"mock docx 2")
    stem2 = resolve_unique_stem("foo", output_dir=tmp_path)
    assert stem2 == "foo_3"


def test_save_artifacts_docx_only(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts saves .docx file in binary mode when OutputFormat.DOCX is configured."""
    config = JobConfig(output_format=OutputFormat.DOCX)
    saved = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)

    assert "docx" in saved
    assert "markdown" not in saved
    assert "json" not in saved

    expected_docx = tmp_path / "financial_audit.docx"
    assert saved["docx"] == expected_docx.resolve()
    assert expected_docx.exists()

    # Reopen document to verify it is valid Word OpenXML
    doc = Document(str(expected_docx))
    assert any("Financial Report 2026" in p.text for p in doc.paragraphs)
    assert any("Expenses Breakdown" in p.text for p in doc.paragraphs)


def test_save_artifacts_docx_atomic_write_and_tmp_cleanup_on_error(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts unlinks .docx.tmp and preserves existing .docx on write error."""
    config = JobConfig(output_format=OutputFormat.DOCX)

    # 1. Clean run writes docx file
    saved = save_artifacts(sample_ocr_result, config, output_dir=tmp_path)
    docx_file = saved["docx"]
    assert docx_file.exists()
    orig_bytes = docx_file.read_bytes()

    # 2. Simulate failure during write_bytes on .docx.tmp
    real_write_bytes = Path.write_bytes

    def failing_write_bytes(self, data, *args, **kwargs):
        if str(self).endswith(".docx.tmp"):
            raise OSError("Simulated disk error during docx temp write")
        return real_write_bytes(self, data, *args, **kwargs)

    with patch.object(Path, "write_bytes", side_effect=failing_write_bytes, autospec=True):
        with pytest.raises(OSError, match="Simulated disk error during docx temp write"):
            save_artifacts(sample_ocr_result, config, output_dir=tmp_path)

    # Verify original file preserved and no lingering .tmp files
    assert docx_file.read_bytes() == orig_bytes
    assert not list(tmp_path.glob("*.tmp"))


def test_save_artifacts_docx_sanitizes_path_leak_on_error(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts sanitizes absolute user paths when DOCX write fails."""
    config = JobConfig(output_format=OutputFormat.DOCX)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_bytes = Path.write_bytes

    def failing_write_bytes(self, data, *args, **kwargs):
        if str(self).endswith(".docx.tmp"):
            # Raise an error containing the raw absolute private path
            raise OSError(f"Permission denied on private file: '{raw_path}'")
        return real_write_bytes(self, data, *args, **kwargs)

    with patch.object(Path, "write_bytes", side_effect=failing_write_bytes, autospec=True):
        with pytest.raises(OSError) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

        err_str = str(exc_info.value)
        # Verify the raw absolute path was sanitized to relativized form and not leaked
        assert str(raw_path) not in err_str
        assert "secret_doc.pdf" in err_str


def test_save_artifacts_docx_sanitizes_real_oserror_with_attributes(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts handles real multi-arg PermissionError without crashing and preserves attributes."""
    config = JobConfig(output_format=OutputFormat.DOCX)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_bytes = Path.write_bytes

    def failing_write_bytes(self, data, *args, **kwargs):
        if str(self).endswith(".docx.tmp"):
            # Standard multi-arg OS exception: (errno, strerror, filename)
            raise PermissionError(13, "Permission denied", str(raw_path))
        return real_write_bytes(self, data, *args, **kwargs)

    with patch.object(Path, "write_bytes", side_effect=failing_write_bytes, autospec=True):
        with pytest.raises(PermissionError) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    exc = exc_info.value
    # (a) Confirm reconstruction did not crash and preserved the exact exception class
    assert isinstance(exc, PermissionError)
    # (b) Confirm OS attributes are preserved and filename is sanitized
    assert exc.errno == 13
    assert exc.filename == "secret_doc.pdf"
    assert str(raw_path) not in (exc.filename or "")
    # (c) Confirm str(exc) does not leak the absolute path
    err_str = str(exc)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str


def test_save_artifacts_docx_sanitizes_custom_multi_arg_exception(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts falls back cleanly without TypeError when exception has non-single-string init."""
    class CustomMultiArgError(Exception):
        def __init__(self, code: int, details: str):
            super().__init__(code, details)
            self.code = code
            self.details = details

    config = JobConfig(output_format=OutputFormat.DOCX)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_bytes = Path.write_bytes

    def failing_write_bytes(self, data, *args, **kwargs):
        if str(self).endswith(".docx.tmp"):
            raise CustomMultiArgError(500, f"Critical storage fault on '{raw_path}'")
        return real_write_bytes(self, data, *args, **kwargs)

    with patch.object(Path, "write_bytes", side_effect=failing_write_bytes, autospec=True):
        with pytest.raises(Exception) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    err_str = str(exc_info.value)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str


def test_format_output_docx_sanitizes_embedded_error_message(tmp_path: Path) -> None:
    """Verify format_output redacts user home directories when rendering an OCRResult with error into DOCX."""
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    error_msg = f"Fatal pipeline error processing '{raw_path}'"
    result = OCRResult(
        file_path=str(raw_path),
        pages=[],
        error=error_msg,
        status=JobStatus.FAILED,
    )

    outputs = format_output(
        result,
        OutputFormat.DOCX,
        sanitize_path=True,
        base_dir=tmp_path / "private_user_dir",
    )

    assert "docx" in outputs
    assert isinstance(outputs["docx"], bytes)

    # Reopen document and inspect error paragraph text
    doc = Document(io.BytesIO(outputs["docx"]))
    doc_text = " ".join(p.text for p in doc.paragraphs)
    assert str(raw_path) not in doc_text
    assert "secret_doc.pdf" in doc_text


def test_save_artifacts_markdown_sanitizes_real_oserror_with_attributes(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts handles real multi-arg PermissionError without crashing and preserves attributes on markdown write."""
    config = JobConfig(output_format=OutputFormat.MARKDOWN)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_text = Path.write_text

    def failing_write_text(self, data, *args, **kwargs):
        if str(self).endswith(".md.tmp"):
            raise PermissionError(13, "Permission denied", str(raw_path))
        return real_write_text(self, data, *args, **kwargs)

    with patch.object(Path, "write_text", side_effect=failing_write_text, autospec=True):
        with pytest.raises(PermissionError) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    exc = exc_info.value
    assert isinstance(exc, PermissionError)
    assert exc.errno == 13
    assert exc.filename == "secret_doc.pdf"
    assert str(raw_path) not in (exc.filename or "")
    err_str = str(exc)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str


def test_save_artifacts_markdown_sanitizes_custom_multi_arg_exception(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts falls back cleanly without TypeError when exception has non-single-string init on markdown write."""
    class CustomMultiArgError(Exception):
        def __init__(self, code: int, details: str):
            super().__init__(code, details)
            self.code = code
            self.details = details

    config = JobConfig(output_format=OutputFormat.MARKDOWN)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_text = Path.write_text

    def failing_write_text(self, data, *args, **kwargs):
        if str(self).endswith(".md.tmp"):
            raise CustomMultiArgError(500, f"Critical storage fault on '{raw_path}'")
        return real_write_text(self, data, *args, **kwargs)

    with patch.object(Path, "write_text", side_effect=failing_write_text, autospec=True):
        with pytest.raises(Exception) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    err_str = str(exc_info.value)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str


def test_save_artifacts_json_sanitizes_real_oserror_with_attributes(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts handles real multi-arg PermissionError without crashing and preserves attributes on json write."""
    config = JobConfig(output_format=OutputFormat.JSON)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_text = Path.write_text

    def failing_write_text(self, data, *args, **kwargs):
        if str(self).endswith(".json.tmp"):
            raise PermissionError(13, "Permission denied", str(raw_path))
        return real_write_text(self, data, *args, **kwargs)

    with patch.object(Path, "write_text", side_effect=failing_write_text, autospec=True):
        with pytest.raises(PermissionError) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    exc = exc_info.value
    assert isinstance(exc, PermissionError)
    assert exc.errno == 13
    assert exc.filename == "secret_doc.pdf"
    assert str(raw_path) not in (exc.filename or "")
    err_str = str(exc)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str


def test_save_artifacts_json_sanitizes_custom_multi_arg_exception(
    sample_ocr_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify save_artifacts falls back cleanly without TypeError when exception has non-single-string init on json write."""
    class CustomMultiArgError(Exception):
        def __init__(self, code: int, details: str):
            super().__init__(code, details)
            self.code = code
            self.details = details

    config = JobConfig(output_format=OutputFormat.JSON)
    raw_path = tmp_path / "private_user_dir" / "secret_doc.pdf"
    sample_ocr_result.file_path = str(raw_path)

    real_write_text = Path.write_text

    def failing_write_text(self, data, *args, **kwargs):
        if str(self).endswith(".json.tmp"):
            raise CustomMultiArgError(500, f"Critical storage fault on '{raw_path}'")
        return real_write_text(self, data, *args, **kwargs)

    with patch.object(Path, "write_text", side_effect=failing_write_text, autospec=True):
        with pytest.raises(Exception) as exc_info:
            save_artifacts(
                sample_ocr_result,
                config,
                output_dir=tmp_path / "out",
                base_dir=tmp_path / "private_user_dir",
                sanitize_path=True,
            )

    err_str = str(exc_info.value)
    assert str(raw_path) not in err_str
    assert "secret_doc.pdf" in err_str






