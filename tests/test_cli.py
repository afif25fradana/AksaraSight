import io
import json
from pathlib import Path
import subprocess
import sys
from unittest.mock import MagicMock, patch
from docx import Document
import pytest

from cli.main import SUPPORTED_EXTENSIONS, build_parser, discover_files, main
from core.models import JobStatus, OCRResult, OutputFormat, PageResult


# ==============================================================================
# Fixtures & Helpers
# ==============================================================================

@pytest.fixture
def dummy_png(tmp_path: Path) -> Path:
    p = tmp_path / "sample.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\nDummyPNG")
    return p


@pytest.fixture
def dummy_pdf(tmp_path: Path) -> Path:
    p = tmp_path / "sample.pdf"
    p.write_bytes(b"%PDF-1.4\nDummyPDF")
    return p


@pytest.fixture
def mock_success_result() -> OCRResult:
    return OCRResult(
        file_path="sample.png",
        pages=[
            PageResult(
                page_num=1,
                markdown="# Sample OCR Output\nRecognized text line.",
                status=JobStatus.SUCCESS,
            )
        ],
        total_duration=0.42,
        status=JobStatus.SUCCESS,
        aborted=False,
    )


# ==============================================================================
# Extension Discovery Tests
# ==============================================================================

def test_discover_files_explicit_whitelist(tmp_path: Path) -> None:
    # Supported
    (tmp_path / "img1.PNG").write_bytes(b"1")
    (tmp_path / "img2.jpg").write_bytes(b"2")
    (tmp_path / "doc.pdf").write_bytes(b"3")
    (tmp_path / "scan.tiff").write_bytes(b"4")
    (tmp_path / "photo.webp").write_bytes(b"5")

    # Unsupported
    (tmp_path / "notes.txt").write_bytes(b"6")
    (tmp_path / "archive.zip").write_bytes(b"7")
    (tmp_path / "report.docx").write_bytes(b"8")

    sub_dir = tmp_path / "subdir"
    sub_dir.mkdir()
    (sub_dir / "sub_doc.pdf").write_bytes(b"9")

    # Non-recursive
    files = discover_files(tmp_path, recursive=False)
    names = [f.name for f in files]
    assert sorted(names) == ["doc.pdf", "img1.PNG", "img2.jpg", "photo.webp", "scan.tiff"]

    # Recursive
    rec_files = discover_files(tmp_path, recursive=True)
    rec_names = [f.name for f in rec_files]
    assert sorted(rec_names) == ["doc.pdf", "img1.PNG", "img2.jpg", "photo.webp", "scan.tiff", "sub_doc.pdf"]


# ==============================================================================
# Single File Execution Tests
# ==============================================================================

@patch("cli.main.OCREngine")
def test_cli_single_file_stdout_markdown(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    capsys: pytest.CaptureFixture[str],
) -> None:
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(dummy_png)])

    assert exit_code == 0
    captured = capsys.readouterr()
    assert "# Sample OCR Output" in captured.out
    assert "Recognized text line." in captured.out
    assert captured.err == ""


@patch("cli.main.OCREngine")
def test_cli_single_file_stdout_json_quiet_pipeable(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    capsys: pytest.CaptureFixture[str],
) -> None:
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(dummy_png), "-f", "json", "--quiet"])

    assert exit_code == 0
    captured = capsys.readouterr()
    # Output must be 100% parseable JSON with 0 prefix/suffix text
    parsed = json.loads(captured.out)
    assert parsed["status"] == "SUCCESS"
    assert parsed["pages"][0]["markdown"] == "# Sample OCR Output\nRecognized text line."
    assert captured.err == ""


@patch("cli.main.OCREngine")
def test_cli_single_file_stdout_json_sanitizes_path(
    mock_engine_cls: MagicMock,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify single-file JSON stdout streaming sanitizes file_path and error strings."""
    doc_file = tmp_path / "user_private" / "invoice.png"
    doc_file.parent.mkdir(parents=True, exist_ok=True)
    doc_file.write_bytes(b"\x89PNG\r\n\x1a\nFakePNG")

    raw_error = f"Error processing '{doc_file}'"
    mock_result = OCRResult(
        file_path=str(doc_file),
        pages=[PageResult(page_num=1, markdown="# Invoice", status=JobStatus.SUCCESS)],
        status=JobStatus.SUCCESS,
        error=raw_error,
    )
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(doc_file), "-f", "json", "--quiet"])
    assert exit_code == 0

    captured = capsys.readouterr()
    parsed = json.loads(captured.out)
    assert parsed["file_path"] != str(doc_file.resolve())
    assert str(doc_file.resolve()) not in parsed["error"]


@patch("cli.main.OCREngine")
def test_cli_single_file_save_to_output_dir(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    tmp_path: Path,
    mock_success_result: OCRResult,
) -> None:
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    out_dir = tmp_path / "out"
    exit_code = main([str(dummy_png), "-o", str(out_dir), "-f", "both"])

    assert exit_code == 0
    assert (out_dir / "sample.md").exists()
    assert (out_dir / "sample.json").exists()


# ==============================================================================
# Fatal Configuration & Validation Tests (Exit Code 1)
# ==============================================================================

def test_cli_fatal_nonexistent_input_path(capsys: pytest.CaptureFixture[str]) -> None:
    exit_code = main(["nonexistent_path_xyz.pdf"])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Error: Input path does not exist" in captured.err


def test_cli_fatal_directory_without_output_flag(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    exit_code = main([str(tmp_path)])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Error: -o/--output directory is required" in captured.err


def test_cli_fatal_directory_with_no_supported_files(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    (tmp_path / "file.txt").write_text("plain text")
    out_dir = tmp_path / "out"
    exit_code = main([str(tmp_path), "-o", str(out_dir)])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Error: No supported document files found" in captured.err


def test_cli_fatal_invalid_backend_override(
    dummy_png: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # argparse should reject invalid backend
    with pytest.raises(SystemExit) as exc:
        main([str(dummy_png), "--backend", "unknown-backend"])
    assert exc.value.code == 2


@patch("cli.main.OCREngine")
def test_cli_fatal_backend_offline_abort(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    mock_engine = MagicMock()
    offline_result = OCRResult(
        file_path=str(dummy_png),
        pages=[PageResult(page_num=1, status=JobStatus.FAILED, error="Connection refused")],
        status=JobStatus.FAILED,
        error="Inference backend offline on page 1: Connection refused",
        aborted=True,
    )
    mock_engine.process_document.return_value = offline_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(dummy_png)])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Aborted: Inference backend offline" in captured.err


# ==============================================================================
# Batch Processing & Exit Code 2 (Partial/Failure) Tests
# ==============================================================================

@patch("cli.main.OCREngine")
def test_cli_batch_folder_all_success(
    mock_engine_cls: MagicMock,
    tmp_path: Path,
    mock_success_result: OCRResult,
) -> None:
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    (docs_dir / "a.png").write_bytes(b"a")
    (docs_dir / "b.pdf").write_bytes(b"b")

    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    out_dir = tmp_path / "results"
    exit_code = main([str(docs_dir), "-o", str(out_dir)])

    assert exit_code == 0
    assert mock_engine.process_document.call_count == 2


@patch("cli.main.OCREngine")
def test_cli_batch_folder_partial_failure_exit_code_2(
    mock_engine_cls: MagicMock,
    tmp_path: Path,
    mock_success_result: OCRResult,
) -> None:
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    (docs_dir / "1.png").write_bytes(b"1")
    (docs_dir / "2.png").write_bytes(b"2")

    failed_result = OCRResult(
        file_path="2.png",
        pages=[PageResult(page_num=1, status=JobStatus.FAILED, error="Timeout")],
        status=JobStatus.FAILED,
        error="Inference failure",
        aborted=False,
    )

    mock_engine = MagicMock()
    mock_engine.process_document.side_effect = [mock_success_result, failed_result]
    mock_engine_cls.return_value = mock_engine

    out_dir = tmp_path / "results"
    exit_code = main([str(docs_dir), "-o", str(out_dir)])

    # Exit code 2 for partial failure
    assert exit_code == 2


# ==============================================================================
# Settings Override & Parameter Passing Test
# ==============================================================================

@patch("cli.main.OCREngine")
def test_cli_passes_settings_and_prompt_overrides(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
) -> None:
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([
        str(dummy_png),
        "--backend", "ollama",
        "--endpoint", "http://127.0.0.1:11434/v1",
        "-p", "table",
        "--prompt", "Custom OCR Table Prompt:",
    ])

    assert exit_code == 0
    mock_engine_cls.assert_called_once()
    _, init_kwargs = mock_engine_cls.call_args
    passed_settings = init_kwargs["settings"]
    assert passed_settings.backend == "ollama"
    assert passed_settings.local_endpoint == "http://127.0.0.1:11434/v1"

    _, proc_kwargs = mock_engine.process_document.call_args
    passed_config = proc_kwargs["config"]
    assert passed_config.prompt_mode == "table"
    assert passed_config.custom_prompt == "Custom OCR Table Prompt:"
    assert passed_config.effective_prompt == "Custom OCR Table Prompt:"


def test_cli_remote_endpoint_rejected_without_flag(
    dummy_png: Path,
    capsys: pytest.CaptureFixture,
) -> None:
    """Verify CLI exits with code 1 and error message if remote endpoint given without --allow-remote."""
    exit_code = main([
        str(dummy_png),
        "--endpoint", "http://192.168.1.100:8080/v1",
    ])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Configuration error: Security violation: Non-loopback endpoint" in captured.err


@patch("cli.main.OCREngine")
def test_cli_remote_endpoint_accepted_with_allow_remote_flag(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    capsys: pytest.CaptureFixture,
) -> None:
    """Verify CLI accepts non-loopback endpoint when --allow-remote flag is specified and warns."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([
        str(dummy_png),
        "--endpoint", "http://192.168.1.100:8080/v1",
        "--allow-remote",
    ])
    assert exit_code == 0
    _, init_kwargs = mock_engine_cls.call_args
    passed_settings = init_kwargs["settings"]
    assert passed_settings.local_endpoint == "http://192.168.1.100:8080/v1"
    assert passed_settings.allow_remote is True
    captured = capsys.readouterr()
    assert "WARNING: Backend endpoint is non-loopback" in captured.err


@patch("cli.main.OCREngine")
def test_cli_batch_duplicate_filenames_no_overwrite(
    mock_engine_cls: MagicMock,
    tmp_path: Path,
) -> None:
    """Verify CLI batch processing with duplicate stems creates unique files instead of overwriting."""
    input_dir = tmp_path / "inputs"
    sub_a = input_dir / "sub_a"
    sub_b = input_dir / "sub_b"
    sub_a.mkdir(parents=True)
    sub_b.mkdir(parents=True)

    file_a = sub_a / "document.pdf"
    file_b = sub_b / "document.pdf"
    file_a.write_bytes(b"%PDF-1.4 dummy a")
    file_b.write_bytes(b"%PDF-1.4 dummy b")

    output_dir = tmp_path / "outputs"

    def fake_process_document(path, config=None):
        return OCRResult(
            file_path=str(path),
            pages=[PageResult(page_num=1, markdown=f"# Output for {path.parent.name}", status=JobStatus.SUCCESS)],
            status=JobStatus.SUCCESS,
        )

    mock_engine = MagicMock()
    mock_engine.process_document.side_effect = fake_process_document
    mock_engine_cls.return_value = mock_engine

    exit_code = main([
        str(input_dir),
        "-r",
        "-o", str(output_dir),
        "-f", "markdown",
    ])

    assert exit_code == 0
    assert (output_dir / "document.md").exists()
    assert (output_dir / "document_2.md").exists()
    content_1 = (output_dir / "document.md").read_text(encoding="utf-8")
    content_2 = (output_dir / "document_2.md").read_text(encoding="utf-8")
    assert content_1 != content_2


@patch("cli.main.OCREngine")
def test_cli_max_pages_flag_forwarded(
    mock_engine_cls: MagicMock,
    dummy_pdf: Path,
    mock_success_result: OCRResult,
) -> None:
    """Verify --max-pages flag forwards positive int to JobConfig."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(dummy_pdf), "--max-pages", "5"])

    assert exit_code == 0
    assert mock_engine.process_document.call_count == 1
    _, kwargs = mock_engine.process_document.call_args
    assert kwargs["config"].max_pages == 5


def test_cli_max_pages_flag_invalid(
    dummy_pdf: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify non-positive --max-pages flag triggers fatal error."""
    exit_code = main([str(dummy_pdf), "--max-pages", "0"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Error: --max-pages must be a positive integer" in captured.err


@patch("cli.main.OCREngine")
def test_cli_dpi_flag_forwarded(
    mock_engine_cls: MagicMock,
    dummy_pdf: Path,
    mock_success_result: OCRResult,
) -> None:
    """Verify --dpi flag forwards positive int to JobConfig."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    exit_code = main([str(dummy_pdf), "--dpi", "150"])

    assert exit_code == 0
    assert mock_engine.process_document.call_count == 1
    _, kwargs = mock_engine.process_document.call_args
    assert kwargs["config"].dpi == 150


def test_cli_dpi_flag_invalid(
    dummy_pdf: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify non-positive --dpi flag triggers fatal error."""
    exit_code = main([str(dummy_pdf), "--dpi", "0"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Error: --dpi must be a positive integer" in captured.err


@patch("cli.main.OCREngine")
def test_cli_batch_streaming_exit_codes(
    mock_engine_cls: MagicMock,
    tmp_path: Path,
    mock_success_result: OCRResult,
) -> None:
    """Verify batch processing correctly resolves exit codes without accumulating results."""
    in_dir = tmp_path / "in"
    out_dir = tmp_path / "out"
    in_dir.mkdir()
    out_dir.mkdir()

    f1 = in_dir / "doc1.png"
    f2 = in_dir / "doc2.png"
    f1.write_bytes(b"data1")
    f2.write_bytes(b"data2")

    mock_engine = MagicMock()
    # 1. All success -> exit code 0
    mock_engine.process_document.side_effect = [
        OCRResult(file_path=str(f1), status=JobStatus.SUCCESS),
        OCRResult(file_path=str(f2), status=JobStatus.SUCCESS),
    ]
    mock_engine_cls.return_value = mock_engine

    exit_0 = main([str(in_dir), "-o", str(out_dir), "-q"])
    assert exit_0 == 0

    # 2. One partial/failed -> exit code 2
    mock_engine.process_document.side_effect = [
        OCRResult(file_path=str(f1), status=JobStatus.SUCCESS),
        OCRResult(file_path=str(f2), status=JobStatus.FAILED),
    ]
    exit_2 = main([str(in_dir), "-o", str(out_dir), "-q"])
    assert exit_2 == 2

    # 3. One aborted -> exit code 1
    mock_engine.process_document.side_effect = [
        OCRResult(file_path=str(f1), status=JobStatus.SUCCESS),
        OCRResult(file_path=str(f2), status=JobStatus.FAILED, aborted=True),
    ]
    exit_1 = main([str(in_dir), "-o", str(out_dir), "-q"])
    assert exit_1 == 1


# ==============================================================================
# Hardware Detection CLI Flag Tests
# ==============================================================================

@patch("cli.main.detect_hardware")
def test_cli_detect_hardware_flag(
    mock_detect: MagicMock,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify --detect-hardware runs detection, prints summary, and exits with 0."""
    from core.hardware import HardwareProfile

    mock_detect.return_value = HardwareProfile(
        gpu_name="NVIDIA Test GPU",
        vram_mb=8192,
        cuda_available=True,
        cuda_supported=True,
        vulkan_available=True,
        cpu_name="Test CPU",
        recommended_backend="cuda",
        details="Test CUDA details",
    )

    exit_code = main(["--detect-hardware"])
    assert exit_code == 0
    mock_detect.assert_called_once()
    captured = capsys.readouterr()
    assert "SYSTEM HARDWARE DETECTION REPORT" in captured.out
    assert "NVIDIA Test GPU" in captured.out
    assert "RECOMMENDED BACKEND:  CUDA" in captured.out


def test_cli_missing_input_shows_error(
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify omitting input argument without special flags prints error and exits with 1."""
    exit_code = main([])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "the following arguments are required: input" in captured.err


def test_cli_main_launch_subprocess() -> None:
    """Verify 'python -m cli.main --help' launches via __main__ block cleanly."""
    repo_root = Path(__file__).resolve().parent.parent
    result = subprocess.run(
        [sys.executable, "-m", "cli.main", "--help"],
        cwd=repo_root,
        capture_output=True,
        text=True,
        timeout=10,
    )
    assert result.returncode == 0, f"cli.main failed to launch:\n{result.stderr}"
    assert "usage:" in result.stdout.lower()


def test_cli_real_engine_integration(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """Verify CLI main() runs end-to-end with real OCREngine, real Settings, and real pipeline (only VisionClient mocked)."""
    from PIL import Image

    test_img = tmp_path / "invoice.png"
    Image.new("RGB", (150, 150), color="white").save(test_img)

    with patch("core.client.VisionClient.complete", return_value=("# Real Invoice Title\nLine item text", {"id": "test"}, 0.05)):
        exit_code = main([str(test_img), "--dpi", "100", "--max-pages", "1"])

    assert exit_code == 0
    captured = capsys.readouterr()
    assert "# Real Invoice Title" in captured.out
    assert "Line item text" in captured.out


def test_cli_version_flag(capsys: pytest.CaptureFixture[str]) -> None:
    """Verify 'AksaraSight-CLI --version' prints correct version string and exits."""
    from core.constants import __version__

    with pytest.raises(SystemExit) as exc_info:
        main(["--version"])
    assert exc_info.value.code == 0
    captured = capsys.readouterr()
    assert f"AksaraSight-CLI {__version__}" in captured.out


def test_cli_aborts_fast_when_verify_backend_fails(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    """Verify CLI aborts before processing documents if verify_backend fails."""
    from core.client import ClientError

    test_img = tmp_path / "doc.png"
    test_img.write_bytes(b"dummy")

    with patch("cli.main.OCREngine") as mock_engine_cls:
        mock_engine = MagicMock()
        mock_engine.verify_backend.side_effect = ClientError("Backend not responding correctly to image input")
        mock_engine_cls.return_value = mock_engine

        exit_code = main([str(test_img)])

    assert exit_code == 1
    assert mock_engine.process_document.call_count == 0
    captured = capsys.readouterr()
    assert "Error: Backend not responding correctly to image input" in captured.err


# ==============================================================================
# Doctor Flag Tests (--doctor)
# ==============================================================================

@patch("cli.main.OCREngine")
@patch("cli.main.probe_server_health")
@patch("cli.main.get_installed_runtime_path")
@patch("cli.main.is_runtime_installed")
@patch("cli.main.detect_hardware")
def test_cli_doctor_all_pass(
    mock_detect: MagicMock,
    mock_is_installed: MagicMock,
    mock_get_path: MagicMock,
    mock_probe: MagicMock,
    mock_engine_cls: MagicMock,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify --doctor reports healthy status and returns exit code 0 when all checks pass."""
    from core.hardware import HardwareProfile
    from core.server_manager import ServerStatus

    mock_detect.return_value = HardwareProfile(
        gpu_name="NVIDIA RTX 4070",
        vram_mb=8192,
        cuda_available=True,
        cuda_supported=True,
        cuda_driver_version="576.88",
        cpu_name="Test CPU",
        recommended_backend="cuda",
        details="CUDA 12.4 supported",
    )
    mock_is_installed.return_value = True
    mock_get_path.return_value = Path("C:/runtimes/llama-server.exe")
    mock_probe.return_value = (ServerStatus.READY, "Server is healthy and ready")

    mock_engine = MagicMock()
    mock_engine.verify_backend.return_value = None
    mock_engine_cls.return_value = mock_engine

    exit_code = main(["--doctor"])

    assert exit_code == 0
    captured = capsys.readouterr()
    assert "AKSARASIGHT DIAGNOSTIC REPORT" in captured.out
    assert "1. Configuration" in captured.out
    assert "[PASS] Backend:             llama-cpp" in captured.out
    assert "2. Hardware Detection" in captured.out
    assert "[PASS] Primary GPU:         NVIDIA RTX 4070 (8192 MB)" in captured.out
    assert "3. Runtime Installation (Managed Mode)" in captured.out
    assert "[PASS] Managed Runtime:     b10930-cuda (INSTALLED)" in captured.out
    assert "4. Server Reachability" in captured.out
    assert "[PASS] Endpoint Health:     READY (http://localhost:8080/health)" in captured.out
    assert "5. Multimodal Vision Probe" in captured.out
    assert "[PASS] 1x1 Image Test:      VERIFIED (Vision projector active, inference operational)" in captured.out
    assert "STATUS: HEALTHY - All checks passed (5/5)." in captured.out


@patch("cli.main.OCREngine")
@patch("cli.main.probe_server_health")
@patch("cli.main.detect_hardware")
def test_cli_doctor_custom_runtime_mode(
    mock_detect: MagicMock,
    mock_probe: MagicMock,
    mock_engine_cls: MagicMock,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Verify Step 3 reports Custom Path existence and does NOT call is_runtime_installed."""
    from core.hardware import HardwareProfile
    from core.server_manager import ServerStatus

    mock_detect.return_value = HardwareProfile(
        cpu_name="Test CPU",
        recommended_backend="cpu",
        details="CPU inference",
    )
    mock_probe.return_value = (ServerStatus.READY, "Server is healthy and ready")
    mock_engine = MagicMock()
    mock_engine.verify_backend.return_value = None
    mock_engine_cls.return_value = mock_engine

    # Case A: Valid custom binary exists
    custom_exe = tmp_path / "custom-llama-server.exe"
    custom_exe.write_bytes(b"MZ_DUMMY_EXE")

    monkeypatch.setenv("OCR_RUNTIME_MODE", "custom")
    monkeypatch.setenv("OCR_LLAMA_SERVER_PATH", str(custom_exe))

    with patch("cli.main.is_runtime_installed") as mock_managed_check:
        exit_code = main(["--doctor"])
        # Crucial check: is_runtime_installed must NOT be called in custom mode
        mock_managed_check.assert_not_called()

    assert exit_code == 0
    captured = capsys.readouterr()
    assert "3. Runtime Installation (Custom Path)" in captured.out
    assert "[PASS] Custom Binary:       FOUND" in captured.out
    assert str(custom_exe) in captured.out
    assert "STATUS: HEALTHY" in captured.out

    # Case B: Custom binary does NOT exist
    missing_exe = tmp_path / "nonexistent.exe"
    monkeypatch.setenv("OCR_LLAMA_SERVER_PATH", str(missing_exe))

    exit_code = main(["--doctor"])
    assert exit_code == 1
    captured = capsys.readouterr()
    assert "3. Runtime Installation (Custom Path)" in captured.out
    assert "[FAIL] Custom Binary:       NOT FOUND" in captured.out
    assert "STATUS: UNHEALTHY - 1 check failed." in captured.out
    assert f"Verify the custom binary path exists: '{missing_exe}'." in captured.out


@patch("cli.main.probe_server_health")
@patch("cli.main.is_runtime_installed")
@patch("cli.main.detect_hardware")
def test_cli_doctor_server_offline(
    mock_detect: MagicMock,
    mock_is_installed: MagicMock,
    mock_probe: MagicMock,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify --doctor handles offline server, skips multimodal probe, and exits with 1."""
    from core.hardware import HardwareProfile
    from core.server_manager import ServerStatus

    mock_detect.return_value = HardwareProfile(cpu_name="CPU", recommended_backend="cpu")
    mock_is_installed.return_value = True
    mock_probe.return_value = (ServerStatus.OFFLINE, "Connection refused (server not running)")

    with patch("cli.main.get_installed_runtime_path", return_value=Path("C:/llama-server.exe")), \
         patch("cli.main.OCREngine") as mock_engine_cls:
        exit_code = main(["--doctor"])
        mock_engine_cls.assert_not_called()

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "[FAIL] Endpoint Health:     OFFLINE" in captured.out
    assert "Connection refused (server not running)" in captured.out
    assert "[SKIP] 1x1 Image Test:      SKIPPED (Server is not ready)" in captured.out
    assert "STATUS: UNHEALTHY - 1 check failed." in captured.out
    assert "Start the backend server via GUI or run 'llama-server'" in captured.out


@patch("cli.main.probe_server_health")
@patch("cli.main.is_runtime_installed")
@patch("cli.main.detect_hardware")
def test_cli_doctor_runtime_not_installed(
    mock_detect: MagicMock,
    mock_is_installed: MagicMock,
    mock_probe: MagicMock,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify --doctor reports missing managed runtime and exits with 1."""
    from core.hardware import HardwareProfile
    from core.server_manager import ServerStatus

    mock_detect.return_value = HardwareProfile(cpu_name="CPU", recommended_backend="cpu")
    mock_is_installed.return_value = False
    mock_probe.return_value = (ServerStatus.READY, "Server is healthy and ready")

    with patch("cli.main.OCREngine") as mock_engine_cls:
        mock_engine = MagicMock()
        mock_engine.verify_backend.return_value = None
        mock_engine_cls.return_value = mock_engine

        exit_code = main(["--doctor"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "[FAIL] Managed Runtime:     b10930-cpu (NOT INSTALLED)" in captured.out
    assert "STATUS: UNHEALTHY - 1 check failed." in captured.out
    assert "Install the managed runtime 'b10930-cpu' via GUI Settings" in captured.out


@patch("cli.main.OCREngine")
@patch("cli.main.probe_server_health")
@patch("cli.main.is_runtime_installed")
@patch("cli.main.detect_hardware")
def test_cli_doctor_vision_probe_fails(
    mock_detect: MagicMock,
    mock_is_installed: MagicMock,
    mock_probe: MagicMock,
    mock_engine_cls: MagicMock,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify --doctor captures ClientError during multimodal probe and exits with 1."""
    from core.client import ClientError
    from core.hardware import HardwareProfile
    from core.server_manager import ServerStatus

    mock_detect.return_value = HardwareProfile(cpu_name="CPU", recommended_backend="cpu")
    mock_is_installed.return_value = True
    mock_probe.return_value = (ServerStatus.READY, "Server is healthy and ready")

    mock_engine = MagicMock()
    mock_engine.verify_backend.side_effect = ClientError("Missing --mmproj projector")
    mock_engine_cls.return_value = mock_engine

    with patch("cli.main.get_installed_runtime_path", return_value=Path("C:/llama-server.exe")):
        exit_code = main(["--doctor"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "[FAIL] 1x1 Image Test:      FAILED (Missing --mmproj projector)" in captured.out
    assert "STATUS: UNHEALTHY - 1 check failed." in captured.out
    assert "Ensure the backend was launched with multimodal vision projector support (--mmproj)." in captured.out


def test_cli_doctor_config_error(capsys: pytest.CaptureFixture[str]) -> None:
    """Verify --doctor handles non-loopback endpoint security error and exits with 1."""
    exit_code = main(["--doctor", "--endpoint", "http://remote-machine.internal:8080/v1"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "1. Configuration" in captured.out
    assert "[FAIL] Configuration:       Invalid settings:" in captured.out
    assert "Security violation: Non-loopback endpoint" in captured.out
    assert "3. Runtime Installation" in captured.out
    assert "[SKIP] Runtime Check:       SKIPPED (Configuration error)" in captured.out
    assert "STATUS: UNHEALTHY" in captured.out


# ==============================================================================
# DOCX Export CLI Tests
# ==============================================================================

def test_cli_parser_docx_choice() -> None:
    """Verify CLI parser accepts 'docx' as a valid -f/--format choice."""
    parser = build_parser()
    args = parser.parse_args(["sample.pdf", "-f", "docx"])
    assert args.format == "docx"


@patch("cli.main.OCREngine")
def test_cli_single_file_docx_output_dir(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    tmp_path: Path,
) -> None:
    """Verify CLI writes {stem}.docx when -o is provided, readable back by Document()."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    out_dir = tmp_path / "docx_out"
    exit_code = main([str(dummy_png), "-f", "docx", "-o", str(out_dir)])

    assert exit_code == 0
    expected_docx = out_dir / f"{dummy_png.stem}.docx"
    assert expected_docx.exists()

    # Read back and verify valid Document with expected content
    doc = Document(str(expected_docx))
    doc_text = " ".join(p.text for p in doc.paragraphs)
    assert "Sample OCR Output" in doc_text
    assert "Recognized text line." in doc_text


@patch("cli.main.OCREngine")
def test_cli_single_file_docx_stdout_terminal_isatty_refused(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify CLI refuses to dump binary DOCX to a terminal (isatty=True) and exits with 1."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    with patch.object(sys.stdout, "isatty", return_value=True):
        exit_code = main([str(dummy_png), "-f", "docx"])

    assert exit_code == 1
    captured = capsys.readouterr()
    assert "Cannot write binary DOCX output to a terminal" in captured.err
    # Ensure no binary data was written to stdout
    assert captured.out == ""


@patch("cli.main.OCREngine")
def test_cli_single_file_docx_stdout_piped_writes_buffer(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
) -> None:
    """Verify CLI streams raw binary bytes to sys.stdout.buffer when redirected (isatty=False)."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    mock_buffer = io.BytesIO()
    fake_stdout = MagicMock()
    fake_stdout.isatty.return_value = False
    fake_stdout.buffer = mock_buffer

    with patch("sys.stdout", fake_stdout):
        exit_code = main([str(dummy_png), "-f", "docx"])

    assert exit_code == 0
    # Verify write was called on buffer, not sys.stdout.write
    fake_stdout.write.assert_not_called()

    # Verify raw bytes written to buffer represent a valid Word OpenXML document
    binary_data = mock_buffer.getvalue()
    assert len(binary_data) > 0
    doc = Document(io.BytesIO(binary_data))
    doc_text = " ".join(p.text for p in doc.paragraphs)
    assert "Sample OCR Output" in doc_text
    assert "Recognized text line." in doc_text


@patch("cli.main.OCREngine")
def test_cli_existing_formats_untouched(
    mock_engine_cls: MagicMock,
    dummy_png: Path,
    mock_success_result: OCRResult,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Verify markdown, json, and both format outputs retain existing stdout behavior."""
    mock_engine = MagicMock()
    mock_engine.process_document.return_value = mock_success_result
    mock_engine_cls.return_value = mock_engine

    # 1. Markdown
    exit_code = main([str(dummy_png), "-f", "markdown"])
    assert exit_code == 0
    out_md = capsys.readouterr().out
    assert "# Sample OCR Output" in out_md

    # 2. JSON
    exit_code = main([str(dummy_png), "-f", "json", "-q"])
    assert exit_code == 0
    out_json = capsys.readouterr().out
    data = json.loads(out_json)
    assert data["status"] == "SUCCESS"

    # 3. BOTH
    exit_code = main([str(dummy_png), "-f", "both"])
    assert exit_code == 0
    out_both = capsys.readouterr().out
    assert "# Sample OCR Output" in out_both
    assert "\n\n---\n\n" in out_both
    assert '"status": "SUCCESS"' in out_both



